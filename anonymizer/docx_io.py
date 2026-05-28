from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph


@dataclass(frozen=True)
class TextBlock:
    index: int
    text: str


def _iter_paragraphs(container) -> Iterable[Paragraph]:
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def iter_all_paragraphs(doc: DocumentObject) -> Iterable[Paragraph]:
    yield from _iter_paragraphs(doc)
    for section in doc.sections:
        yield from _iter_paragraphs(section.header)
        yield from _iter_paragraphs(section.footer)


def extract_docx_text(path: str | Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in iter_all_paragraphs(doc) if p.text is not None)


def load_blocks(path: str | Path) -> list[TextBlock]:
    doc = Document(str(path))
    return [TextBlock(i, p.text or "") for i, p in enumerate(iter_all_paragraphs(doc))]


def _replace_in_paragraph(paragraph: Paragraph, replacements: list[tuple[str, str]]) -> None:
    original = paragraph.text or ""
    if not original:
        return

    replaced = original
    for value, token in replacements:
        if value:
            replaced = replaced.replace(value, token)
    if replaced == original:
        return

    # Быстрый путь: если значение целиком лежит в одном run, сохраняем run-структуру.
    changed_in_runs = False
    for run in paragraph.runs:
        text = run.text
        new_text = text
        for value, token in replacements:
            if value and value in new_text:
                new_text = new_text.replace(value, token)
        if new_text != text:
            run.text = new_text
            changed_in_runs = True

    if paragraph.text == replaced:
        return

    # Fallback: значение было разрезано между runs. Сохраняем стиль первого run и абзаца,
    # но схлопываем текст. Это осознанный плавный переход для DOCX с дроблеными run'ами.
    if paragraph.runs:
        first = paragraph.runs[0]
        for run in list(paragraph.runs)[1:]:
            run._element.getparent().remove(run._element)
        first.text = replaced
    else:
        paragraph.add_run(replaced)


def anonymize_docx_by_values(input_path: str | Path, output_path: str | Path, value_to_token: dict[str, str]) -> None:
    doc = Document(str(input_path))
    replacements = sorted(value_to_token.items(), key=lambda item: len(item[0]), reverse=True)
    for paragraph in iter_all_paragraphs(doc):
        _replace_in_paragraph(paragraph, replacements)
    doc.core_properties.comments = "Anonymized by regex + LLM pipeline."
    doc.save(str(output_path))
